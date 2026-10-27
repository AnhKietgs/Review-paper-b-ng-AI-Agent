"""Generate original, fictional documents for repeatable local demos."""

from pathlib import Path

import pymupdf
from docx import Document


def generate():
    """Write small text-layer PDF and DOCX with known source positions."""
    root = Path(__file__).parent
    paper = pymupdf.open()
    pages = [
        "Evidence-Grounded Reading Support: A Fictional Pilot Study\n\n"
        "This synthetic paper is authored for the Paper Review Lab demo.\n"
        "Objective: evaluate whether source-linked notes help readers inspect scientific claims.\n"
        "Method: 24 participants read the same two short texts in a within-subject study.\n"
        "Participants received either plain summaries or summaries with source locations.\n"
        "The allocation order was counterbalanced, but participants were not blinded.",
        "Results and Limitations\n\n"
        "Participants located supporting passages in a median of 42 seconds with source links,\n"
        "compared with 67 seconds with plain summaries.\n"
        "The study reports descriptive statistics only and does not establish causation.\n"
        "Limitations include a small convenience sample and only two reading tasks.\n"
        "The authors claim a practical interface contribution, not universal scientific novelty.",
    ]
    for text in pages:
        page = paper.new_page()
        page.insert_textbox(pymupdf.Rect(54, 60, 540, 760), text, fontsize=11)
    paper.save(root / "main_sample.pdf")
    related = Document()
    related.add_heading("Reading Assistance: Fictional Related Study", level=1)
    related.add_paragraph("This original synthetic document is provided for a local software demo.")
    related.add_paragraph("Method: 40 participants used a reading assistant with highlighted source paragraphs.")
    related.add_paragraph("The study measured task completion time but did not measure long-term comprehension.")
    related.save(root / "related_sample.docx")
    scanned = pymupdf.open()
    scanned.new_page()
    scanned.save(root / "scanned_sample.pdf")


if __name__ == "__main__":
    generate()
